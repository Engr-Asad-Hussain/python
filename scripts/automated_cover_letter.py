from docx import Document
from docx2pdf import convert

def main():
    source_docx = r'C:/Users/Asad-Hussain/Documents/Resume/Cover Letter/Cover Letter - Template.docx'
    destination_docx = r'C:/Users/Asad-Hussain/Documents/Resume/template.docx'
    destination_pdf = r'C:/Users/Asad-Hussain/Documents/Cover Letter - Asad Hussain.pdf'

    company = "Axx"
    replacements = {
        "<Today_Date>": "February 15, 2024",
        "<Company_Title>": company,
        "<Company_Location>": "Karachi, Pakistan",
        "<Company>": company
    }

    doc = Document(source_docx)

    # Replace placeholders with actual information
    for old_text, new_text in replacements.items():
        for paragraph in doc.paragraphs:
            if old_text in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace(old_text, new_text)

    # Save the modified Word document
    doc.save(destination_docx)

    # Create a PDF with the modified content
    convert(destination_docx, destination_pdf)
    
if __name__ == '__main__':
    main()